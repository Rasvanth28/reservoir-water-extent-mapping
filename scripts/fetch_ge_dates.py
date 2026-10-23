"""List the Google Earth historical imagery dates over each validation point (for the labeling app).

Runs GEHistoricalImagery `availability` (provider TM, zoom 19 = full history; zoom 20 only returns
the latest image) once over the bounding box of a sheet's points, then keeps for every point the
image dates whose coverage touches its 30 m square. Writes `<sheet dir>/<SITE>_ge_dates.json` =
{point_id: ["YYYY-MM-DD", ...]}. Only date metadata is queried, never imagery (a few seconds per site).

The tool is not in the repo: download the macOS release from
https://github.com/Mbucari/GEHistoricalImagery/releases into tools/GEHistoricalImagery/.
Checked on 70 labeled PIL points: every GE date the interpreter used was in the list, and the
bounding-box result equals a per-point query for all 70.

Usage: .venv/bin/python scripts/fetch_ge_dates.py data/ground_truth/full/PIL/PIL_labels.csv [...]
(.csv or .xlsx; only point_id/lon/lat are read)
"""

import argparse
import csv
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TOOL_DIR = ROOT / "tools" / "GEHistoricalImagery"
HALF_DEG = 0.000135  # ~15 m
ZOOM = 19


def points(sheet: Path) -> list[tuple[str, float, float]]:
    if sheet.suffix == ".xlsx":
        rows = list(load_workbook(sheet, read_only=True).worksheets[0].iter_rows(values_only=True))
        recs = [dict(zip(rows[0], r)) for r in rows[1:] if r and r[0]]
    else:
        with open(sheet, newline="") as f:
            recs = list(csv.DictReader(f))
    return [(r["point_id"], float(r["lon"]), float(r["lat"])) for r in recs]


def site_dates(pts: list[tuple[str, float, float]]) -> dict[str, list[str]]:
    """One availability query over the bounding box of all points; per point, the dates whose
    coverage intersects its 30 m square. (A KML region file only uses its first polygon.)"""
    from shapely.geometry import box, shape
    lons, lats = [p[1] for p in pts], [p[2] for p in pts]
    pad = 2 * HALF_DEG
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "availability.json"
        cmd = [str(TOOL_DIR / "GEHistoricalImagery"), "availability",
               "--lower-left", f"{min(lats) - pad},{min(lons) - pad}",
               "--upper-right", f"{max(lats) + pad},{max(lons) + pad}",
               "--zoom", str(ZOOM), "-o", str(out), "-p", "4", "-q"]
        env = dict(os.environ, GEHistoricalImagery_Cache=str(ROOT / "tools" / "gehi_cache"))
        # the tool ends with an interactive date menu that fails without a console; the JSON is written first
        subprocess.run(cmd, cwd=TOOL_DIR, env=env, capture_output=True, stdin=subprocess.DEVNULL, timeout=1800)
        if not out.exists():
            raise SystemExit("GEHistoricalImagery wrote no availability file")
        feats = json.loads(out.read_text())["features"]
    cover = [(f["properties"]["image_date"], shape(f["geometry"])) for f in feats]
    res = {}
    for pid, lon, lat in pts:
        sq = box(lon - HALF_DEG, lat - HALF_DEG, lon + HALF_DEG, lat + HALF_DEG)
        res[pid] = sorted({d for d, g in cover if g.intersects(sq)})
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sheets", type=Path, nargs="+")
    args = ap.parse_args()
    if not (TOOL_DIR / "GEHistoricalImagery").exists():
        sys.exit(f"GEHistoricalImagery not found in {TOOL_DIR}")
    for sheet in args.sheets:
        pts = points(sheet)
        site = pts[0][0].split("-")[0]
        out = sheet.parent / f"{site}_ge_dates.json"
        res = site_dates(pts)
        out.write_text(json.dumps(res, indent=0, sort_keys=True))
        n = sorted(len(v) for v in res.values())
        print(f"{site}: {len(res)} points -> {out} (dates per point: min {n[0]}, median {n[len(n) // 2]}, max {n[-1]})")


if __name__ == "__main__":
    main()
