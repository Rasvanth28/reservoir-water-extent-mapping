"""Labeling workbook: build an .xlsx from a blind `<SITE>_labels.csv`, and export it back to CSV.

build   Writes `<sheet>.xlsx` next to the CSV with dropdowns (label, confidence, mixed, ref_source)
        and grey helper columns holding the reference chip's sensor and image date. `ref_date_t1/_t2`
        are formulas: typing S2 or LS in `ref_source` fills in the chip's image date when it matches
        the chip's sensor. For GE (or a Copernicus image other than the chip) type the date over the
        formula. Chip info comes from `reference_chips/<SITE>_manifest.json` (sensor/date only, blind).
export  Reads the workbook (formulas are recomputed here, so cached values are not needed) and
        writes `<sheet>_filled.csv` with the original columns, ready for `scripts/pilot_qc.py`.

Usage:
  .venv/bin/python scripts/label_workbook.py build data/ground_truth/full/BHA/BHA_labels.csv
  .venv/bin/python scripts/label_workbook.py export data/ground_truth/full/BHA/BHA_labels.xlsx
In Numbers, export with File > Export To > Excel before running `export`.
"""

import argparse
import csv
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = ["point_id", "site", "lon", "lat", "date_t1", "date_t2",
           "label_t1", "confidence_t1", "mixed_t1", "ref_source_t1", "ref_date_t1",
           "label_t2", "confidence_t2", "mixed_t2", "ref_source_t2", "ref_date_t2",
           "interpreter", "minutes", "notes"]
HELPERS = ["chip_source_t1", "chip_date_t1", "chip_source_t2", "chip_date_t2"]
LISTS = {"label": "Water,Forest,Bare,Other,NA", "confidence": "0,1,2,3",
         "mixed": "TRUE,FALSE", "ref_source": "GE,S2,LS"}
WIDTHS = {"point_id": 11, "site": 12, "lon": 9, "lat": 9, "notes": 40}


def col(name):
    return get_column_letter((COLUMNS + HELPERS).index(name) + 1)


def find_manifest(sheet: Path, site: str) -> Path:
    for d in (sheet.parent, sheet.parent.parent / site):
        p = d / "reference_chips" / f"{site}_manifest.json"
        if p.exists():
            return p
    raise FileNotFoundError(f"no {site}_manifest.json near {sheet}; run scripts/make_reference_chips.py first")


def chip_info(manifest: dict, pid: str, t: str) -> tuple[str, str]:
    m = manifest.get(f"{pid}_{t}") or {}
    if m.get("status") != "ok":
        return "", ""
    return m.get("sensor", ""), m.get("img_date", "")


def ref_date_formula(t: str, row: int) -> str:
    src, csrc, cdate = col(f"ref_source_{t}"), col(f"chip_source_{t}"), col(f"chip_date_{t}")
    return f'=IF({src}{row}="","",IF(UPPER({src}{row})={csrc}{row},{cdate}{row},""))'


def read_table(path: Path) -> list[dict]:
    """Rows of a labeling sheet saved as .csv, .xlsx or .numbers, values formatted as CSV text."""
    if path.suffix == ".csv":
        with open(path, newline="") as f:
            return list(csv.DictReader(f))
    if path.suffix == ".xlsx":
        rows = list(load_workbook(path).worksheets[0].iter_rows(values_only=True))
    elif path.suffix == ".numbers":
        from numbers_parser import Document
        rows = Document(str(path)).sheets[0].tables[0].rows(values_only=True)
    else:
        raise SystemExit(f"unsupported sheet format: {path}")
    header = [str(h).strip() for h in rows[0]]
    return [dict(zip(header, map(fmt, r))) for r in rows[1:] if r and r[0] not in (None, "")]


def merge_labels(rows: list[dict], labels_path: Path) -> int:
    """Copy the labeler's columns (label_t1 ... notes) from another sheet into rows, by point_id."""
    src = {r["point_id"]: r for r in read_table(labels_path)}
    n = 0
    for r in rows:
        s = src.get(r["point_id"])
        if s is None:
            continue
        if (s.get("date_t1"), s.get("date_t2")) != (r["date_t1"], r["date_t2"]):
            raise SystemExit(f"{r['point_id']}: dates differ between {labels_path.name} and the sheet")
        for c in COLUMNS[COLUMNS.index("label_t1"):]:
            r[c] = s.get(c, "")
        n += any(r[c] for c in COLUMNS[COLUMNS.index("label_t1"):])
    return n


def cell_value(c: str, v: str):
    if not v:
        return v
    if c in ("lon", "lat"):
        return float(v)
    if c.startswith("confidence_") or c == "minutes":
        try:
            return int(float(v))
        except ValueError:
            return v
    if c.startswith("mixed_") and v.upper() in ("TRUE", "FALSE"):
        return v.upper() == "TRUE"
    return v


def build(csv_path: Path, force: bool, labels: Path | None = None):
    out = csv_path.with_suffix(".xlsx")
    if out.exists() and not force:
        raise SystemExit(f"{out} exists (it may hold labels); pass --force to overwrite")
    with open(csv_path, newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit(f"{csv_path} is empty")
    if labels:
        print(f"carried over labels for {merge_labels(rows, labels)} points from {labels}")
    site = rows[0]["point_id"].split("-")[0]
    manifest = json.loads(find_manifest(csv_path, site).read_text())

    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(COLUMNS + HELPERS)
    grey = PatternFill("solid", fgColor="E7E6E6")
    for i, r in enumerate(rows, start=2):
        vals = []
        for c in COLUMNS:
            v = cell_value(c, r.get(c, ""))
            if c.startswith("ref_date_") and not v:
                v = ref_date_formula(c[-2:], i)
            vals.append(v)
        for t in ("t1", "t2"):
            vals.extend(chip_info(manifest, r["point_id"], t))
        ws.append(vals)
    n = len(rows) + 1

    for h in HELPERS:
        for cell in ws[col(h)]:
            cell.fill = grey
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for name in COLUMNS + HELPERS:
        ws.column_dimensions[col(name)].width = WIDTHS.get(name, 13)
    # dates as text so Excel/Numbers don't reformat them
    for name in ("date_t1", "date_t2", "chip_date_t1", "chip_date_t2"):
        for cell in ws[col(name)][1:]:
            cell.number_format = "@"
    for name in ("ref_date_t1", "ref_date_t2"):
        for cell in ws[col(name)][1:]:
            cell.number_format = "yyyy-mm-dd"
    for key, values in LISTS.items():
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        dv.error, dv.errorTitle = f"Use one of: {values}", "Invalid value"
        ws.add_data_validation(dv)
        for t in ("t1", "t2"):
            dv.add(f"{col(f'{key}_{t}')}2:{col(f'{key}_{t}')}{n}")
    ws.freeze_panes = "B2"
    wb.save(out)
    print(f"wrote {out} ({len(rows)} points, site {site})")


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def export(xlsx_path: Path):
    ws = load_workbook(xlsx_path)["labels"]
    header = [c.value for c in ws[1]]
    missing = [c for c in COLUMNS + HELPERS if c not in header]
    if missing:
        raise SystemExit(f"{xlsx_path}: missing columns {missing}")
    idx = {h: i for i, h in enumerate(header)}
    out_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[idx["point_id"]] is None:
            continue
        rec = {c: fmt(row[idx[c]]) for c in COLUMNS}
        for t in ("t1", "t2"):
            rec[f"ref_source_{t}"] = src = rec[f"ref_source_{t}"].upper()
            if rec[f"ref_date_{t}"].startswith("="):  # recompute the formula
                same = src and src == fmt(row[idx[f"chip_source_{t}"]])
                rec[f"ref_date_{t}"] = fmt(row[idx[f"chip_date_{t}"]]) if same else ""
        rec["lon"], rec["lat"] = (fmt(row[idx[c]]) for c in ("lon", "lat"))
        out_rows.append(rec)
    out = xlsx_path.with_name(xlsx_path.stem + "_filled.csv")
    with open(out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(out_rows)
    done = sum(1 for r in out_rows if r["label_t1"] and r["label_t2"])
    print(f"wrote {out} ({done}/{len(out_rows)} points labeled at both dates)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("csv", type=Path, nargs="+")
    b.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    b.add_argument("--labels", type=Path, help="carry over labels already entered in this .csv/.xlsx/.numbers "
                                               "sheet (matched by point_id); use with a single CSV")
    e = sub.add_parser("export")
    e.add_argument("xlsx", type=Path, nargs="+")
    args = ap.parse_args()
    if args.cmd == "build":
        if args.labels and len(args.csv) > 1:
            raise SystemExit("--labels works with one CSV at a time")
        for p in args.csv:
            build(p, args.force, args.labels)
    else:
        for p in args.xlsx:
            export(p)


if __name__ == "__main__":
    main()
